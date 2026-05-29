from pathlib import Path

import pytest
from docx import Document

from services.replacements import ReplacementParseError, extract_docx_text, parse_replacements, render_docx, safe_filename


def test_parse_replacements_supports_common_separators():
    replacements = parse_replacements(
        """
        ООО Старое => ООО Новое
        100 рублей -> 200 рублей
        город = Москва
        дата: 29.05.2026
        """
    )

    assert replacements == {
        "ООО Старое": "ООО Новое",
        "100 рублей": "200 рублей",
        "город": "Москва",
        "дата": "29.05.2026",
    }


def test_parse_replacements_rejects_bad_lines():
    with pytest.raises(ReplacementParseError):
        parse_replacements("строка без разделителя")


def test_render_docx_replaces_paragraphs_tables_headers_and_footers(tmp_path: Path):
    template_path = tmp_path / "template.docx"
    output_path = tmp_path / "output.docx"

    document = Document()
    document.add_paragraph("Договор № OLD-NUMBER")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Сумма OLD-AMOUNT"
    section = document.sections[0]
    section.header.paragraphs[0].text = "HEADER-OLD"
    section.footer.paragraphs[0].text = "FOOTER-OLD"
    document.save(template_path)

    result = render_docx(
        template_path,
        output_path,
        {
            "OLD-NUMBER": "15/2026",
            "OLD-AMOUNT": "150 000 рублей",
            "HEADER-OLD": "HEADER-NEW",
            "FOOTER-OLD": "FOOTER-NEW",
        },
    )

    assert result.replacements_made == 4
    assert result.missing_sources == ()

    rendered = Document(output_path)
    assert "Договор № 15/2026" in [paragraph.text for paragraph in rendered.paragraphs]
    assert rendered.tables[0].cell(0, 0).text == "Сумма 150 000 рублей"
    assert rendered.sections[0].header.paragraphs[0].text == "HEADER-NEW"
    assert rendered.sections[0].footer.paragraphs[0].text == "FOOTER-NEW"


def test_safe_filename_keeps_cyrillic_and_removes_spaces():
    assert safe_filename("Договор Подкаст 1/2") == "Договор_Подкаст_1_2"


def test_extract_docx_text_reads_paragraphs(tmp_path: Path):
    template_path = tmp_path / "template.docx"
    document = Document()
    document.add_paragraph("Первый абзац")
    document.add_paragraph("Второй абзац")
    document.save(template_path)

    assert extract_docx_text(template_path) == "Первый абзац\nВторой абзац"
